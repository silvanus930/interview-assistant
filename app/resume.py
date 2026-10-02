import io
import re
from pathlib import Path

TEXT_EXTENSIONS = {".txt", ".md", ".markdown", ".text"}
SUPPORTED_EXTENSIONS = {".pdf", ".docx", *TEXT_EXTENSIONS}


class ResumeError(Exception):
    pass


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n\n".join((page.extract_text() or "") for page in reader.pages)


def _docx_text(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    # Résumés are often laid out in tables, which paragraphs alone miss.
    for table in doc.tables:
        for row in table.rows:
            cells = []
            for cell in row.cells:
                text = cell.text.strip()
                if text and text not in cells:  # merged cells repeat their text
                    cells.append(text)
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ResumeError(f"Unsupported file type '{ext or filename}'. Use PDF, DOCX, TXT or MD.")
    try:
        if ext == ".pdf":
            text = _pdf_text(data)
        elif ext == ".docx":
            text = _docx_text(data)
        else:
            text = data.decode("utf-8-sig", errors="replace")
    except ResumeError:
        raise
    except Exception as exc:
        raise ResumeError(f"Could not read {filename}: {exc}") from exc

    text = text.replace("\r\n", "\n").replace("\x00", "")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        raise ResumeError(
            "No text found in this file. If it is a scanned PDF, paste the résumé text in manually."
        )
    return text
